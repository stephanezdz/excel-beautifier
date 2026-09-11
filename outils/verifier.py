#!/usr/bin/env python3
"""
Le filet de sécurité d'Excel Beautifier.

Une seule commande, huit contrôles, une réponse claire : ça passe ou ça ne
passe pas. Chaque contrôle correspond à une panne qui est VRAIMENT arrivée
sur ce projet — rien n'est là par principe.

    python3 outils/verifier.py

Il ne modifie rien. Il lit, il exécute dans un coin, il range derrière lui.
"""

import io
import os
import re
import subprocess
import sys
import tempfile
import types

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(RACINE)

VERT, ROUGE, JAUNE, GRIS, NET = "\033[32m", "\033[31m", "\033[33m", "\033[90m", "\033[0m"
if not sys.stdout.isatty() or os.environ.get("NO_COLOR"):
    VERT = ROUGE = JAUNE = GRIS = NET = ""

echecs = []
avertissements = []


def controle(titre):
    """Décorateur : exécute le contrôle et affiche son verdict."""
    def enrobe(fonction):
        try:
            detail = fonction()
        except Exception as erreur:
            echecs.append((titre, f"{type(erreur).__name__} : {erreur}"))
            print(f"  {ROUGE}✕{NET} {titre}")
            print(f"      {ROUGE}{type(erreur).__name__} : {erreur}{NET}")
            return
        if detail is None:
            print(f"  {VERT}✓{NET} {titre}")
        elif isinstance(detail, str):
            print(f"  {VERT}✓{NET} {titre}  {GRIS}{detail}{NET}")
        else:
            echecs.append((titre, detail[1]))
            print(f"  {ROUGE}✕{NET} {titre}")
            for ligne in detail[1].splitlines():
                print(f"      {ROUGE}{ligne}{NET}")
        return fonction
    return enrobe


def git(*args):
    return subprocess.run(["git", *args], capture_output=True, text=True).stdout.strip()


def git_brut(*args):
    """Sans .strip() : git status --porcelain commence par une espace signifiante."""
    return subprocess.run(["git", *args], capture_output=True, text=True).stdout


def app_charge():
    """Rend le module app, ou explique pourquoi il n'est pas disponible."""
    module = globals().get("app")
    if module is None:
        raise RuntimeError("app.py ne s'est pas chargé — voir le contrôle précédent")
    return module


# ─────────────────────────────────────────────────────────────────────
#  Le faux Streamlit : il permet de charger app.py sans lancer de serveur.
# ─────────────────────────────────────────────────────────────────────

class _Zone:
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __getattr__(self, n): return lambda *a, **k: None


class _Etat(dict):
    def __getattr__(self, n):
        try:
            return self[n]
        except KeyError:
            raise AttributeError(n)

    def __setattr__(self, n, v):
        self[n] = v


def poser_faux_streamlit():
    st = types.ModuleType("streamlit")
    for nom in ("set_page_config", "title", "markdown", "header", "divider", "success",
                "subheader", "dataframe", "info", "error", "warning", "caption",
                "rerun", "download_button", "write", "image", "code", "stop"):
        setattr(st, nom, lambda *a, **k: None)
    st.file_uploader = lambda *a, **k: None          # aucun fichier : le corps de l'app est ignoré
    st.button = lambda *a, **k: False
    st.radio = lambda label, options, index=0, **k: options[index]
    st.selectbox = lambda label, options, index=0, **k: options[index]
    st.slider = lambda label, mini, maxi, defaut, *a, **k: defaut
    st.checkbox = lambda label, value=False, **k: value
    st.sidebar = _Zone()
    st.expander = lambda *a, **k: _Zone()
    st.container = lambda *a, **k: _Zone()
    st.columns = lambda n, **k: [_Zone() for _ in range(n if isinstance(n, int) else len(n))]
    st.session_state = _Etat()
    st.cache_data = lambda *a, **k: (a[0] if a and callable(a[0]) else (lambda f: f))
    st.cache_resource = st.cache_data
    sys.modules["streamlit"] = st
    return st


# ─────────────────────────────────────────────────────────────────────
print()
print("─" * 62)
print(" Excel Beautifier — contrôle avant envoi")
print("─" * 62)
print()


@controle("app.py est syntaxiquement correct")
def _():
    import py_compile
    try:
        py_compile.compile("app.py", doraise=True, cfile=tempfile.mktemp())
    except py_compile.PyCompileError as e:
        return False, str(e).strip()


@controle("app.py se charge sans erreur")
def _():
    # 🪤 C'est ce contrôle qui aurait attrapé le bug du 09/09 : une fonction
    #    appelée ligne 79 mais définie ligne 93. Streamlit rejoue le script
    #    à chaque clic, donc l'appel arrivait avant la définition.
    poser_faux_streamlit()
    sys.path.insert(0, RACINE)
    for parasite in ("app",):
        sys.modules.pop(parasite, None)
    import app                                        # noqa: F401
    globals()["app"] = app


@controle("le moteur produit un vrai fichier Excel, en-tête en gras comprise")
def _():
    # 🪤 Ce contrôle attrape les deux autres bugs du 09/09 : le classeur lu en
    #    octets bruts au lieu d'un BytesIO, et le flux vidé par l'aperçu qui
    #    rendait un fichier de 0 octet au téléchargement.
    import openpyxl
    app = app_charge()
    octets = open("test_data.xlsx", "rb").read()
    testes = 0
    for niveau, bloc in app.NIVEAUX.items():
        for index in range(len(bloc["palettes"])):
            palette = app.palette_du_niveau(niveau, index)
            sortie = app.embellir(octets, niveau, palette, 14)
            donnees = sortie.getvalue()
            if len(donnees) < 2000:
                return False, f"{niveau} / {palette['nom']} : fichier de {len(donnees)} octets"
            feuille = openpyxl.load_workbook(io.BytesIO(donnees)).active
            if feuille.max_row < 2 or feuille.max_column < 1:
                return False, f"{niveau} / {palette['nom']} : feuille vide"
            if feuille.freeze_panes != "A2":
                return False, f"{niveau} / {palette['nom']} : l'en-tête n'est pas figé"
            # 🪤 Le gras se relit DANS LE FICHIER, jamais dans le code. Le 11/09,
            #    le modèle local a écrit « cellule.font = cellule.font.bold » 3 fois
            #    sur 6 : aucune erreur, fichier enregistré, rien en gras.
            entetes = [c for c in feuille[1] if c.value is not None]
            if not entetes:
                return False, f"{niveau} / {palette['nom']} : en-tête vide"
            fades = [c.coordinate for c in entetes if not (c.font and c.font.bold)]
            if fades:
                return False, (f"{niveau} / {palette['nom']} : l'en-tête n'est PAS en gras "
                               f"({', '.join(fades[:6])}). Le fichier s'enregistre quand même : "
                               "c'est une panne muette, elle ne se voit qu'ici.")
            testes += 1
    return f"{testes} combinaisons"


@controle("l'aperçu se dessine, accents compris")
def _():
    # 🪤 La police de secours de Pillow ne connaît pas les accents : « Quantité »
    #    y devient « Quantit▯ ». Sans police correcte, l'aperçu ment.
    app = app_charge()
    reguliere, _ = app._police_trouvee()
    if not reguliere:
        return False, ("aucune police capable d'écrire le français n'a été trouvée. "
                       "Sur un serveur Linux, packages.txt doit contenir fonts-dejavu-core.")
    octets = open("test_data.xlsx", "rb").read()
    palette = app.palette_du_niveau("Professionnel", 0)
    image, _, _ = app.apercu_image(app.embellir(octets, "Professionnel", palette, 14).getvalue())
    if image.width < 100 or image.height < 40:
        return False, f"image anormale : {image.size}"
    return os.path.basename(reguliere)


@controle("le fichier de référence test_data.xlsx est intact")
def _():
    # 🪤 Le 09/09, Hermold a modifié le fichier de référence pour s'en servir
    #    de brouillon. L'étalon de mesure ne doit jamais bouger.
    if not git("ls-files", "test_data.xlsx"):
        return False, "test_data.xlsx n'est plus suivi par git"
    modifie = subprocess.run(["git", "diff", "--quiet", "--", "test_data.xlsx"]).returncode
    indexe = subprocess.run(["git", "diff", "--cached", "--quiet", "--", "test_data.xlsx"]).returncode
    if modifie or indexe:
        return False, ("il a été modifié. Pour le remettre en état :\n"
                       "git checkout -- test_data.xlsx")


@controle("aucun fichier de brouillon abandonné à la racine")
def _():
    # 🪤 Le 09/09 encore : original_test.xlsx et output_step1.xlsx laissés
    #    traîner après une session de travail.
    autorises = {"test_data.xlsx"}
    trouves = []
    for ligne in git_brut("status", "--porcelain").splitlines():
        trouve = re.match(r"^(..)\s(.*)$", ligne)
        if not trouve:
            continue
        chemin = trouve.group(2).strip().strip('"')
        if "->" in chemin:                       # renommage : on garde la destination
            chemin = chemin.split("->")[-1].strip()
        if "/" in chemin:
            continue
        if chemin.endswith((".xlsx", ".xls", ".csv", ".png", ".tmp")) and chemin not in autorises:
            trouves.append(chemin)
    if trouves:
        return False, "à supprimer ou à ranger : " + ", ".join(trouves)


@controle("aucun style Excel écrit de travers")
def _():
    # 🪤 « cellule.font = cellule.font.bold » ne lève aucune erreur, s'enregistre
    #    sans rien signaler, et ne met rien en gras. C'est la faute que le modèle
    #    local fait une fois sur deux (mesuré au banc d'essai le 11/09).
    #    openpyxl refuse de modifier un style en place : il faut en fabriquer un
    #    neuf — Font(bold=True, ...), jamais partir de celui de la cellule.
    import ast
    STYLES = {"font", "fill", "border", "alignment"}

    def lit_un_style(noeud):
        """Vrai si l'expression va chercher le style d'une cellule."""
        while isinstance(noeud, ast.Attribute):
            if noeud.attr in STYLES:
                return True
            noeud = noeud.value
        return False

    fautes = []
    arbre = ast.parse(open("app.py", encoding="utf-8").read())
    for noeud in ast.walk(arbre):
        if not isinstance(noeud, ast.Assign):
            continue
        for cible in noeud.targets:
            if not isinstance(cible, ast.Attribute):
                continue
            # cellule.font.bold = True  ->  openpyxl l'interdit
            if isinstance(cible.value, ast.Attribute) and cible.value.attr in STYLES:
                fautes.append((noeud.lineno,
                               f"on modifie {cible.value.attr}.{cible.attr} en place ; "
                               "il faut réaffecter un style neuf"))
                continue
            if cible.attr not in STYLES:
                continue
            # cellule.font = cellule.font.bold  ->  s'enregistre, ne fait rien
            if lit_un_style(noeud.value):
                fautes.append((noeud.lineno,
                               f"le nouveau {cible.attr} est lu sur un style existant ; "
                               "il faut en fabriquer un neuf, par exemple Font(bold=True)"))
            elif isinstance(noeud.value, ast.Constant):
                fautes.append((noeud.lineno,
                               f"{cible.attr} reçoit {noeud.value.value!r} au lieu d'un style"))
    if fautes:
        return False, "\n".join(f"app.py ligne {l} : {quoi}" for l, quoi in fautes)
    return "relu ligne à ligne"


@controle("requirements.txt couvre tout ce qu'app.py importe")
def _():
    # 🪤 Un import ajouté sans la ligne correspondante marche en local et
    #    plante au déploiement, plusieurs minutes plus tard.
    equivalences = {"PIL": "pillow", "openpyxl": "openpyxl", "pandas": "pandas",
                    "streamlit": "streamlit", "yaml": "pyyaml"}
    source = open("app.py", encoding="utf-8").read()
    importes = set(re.findall(r"^\s*(?:import|from)\s+([A-Za-z_][A-Za-z0-9_]*)", source, re.M))
    importes -= set(sys.stdlib_module_names)
    declares = open("requirements.txt", encoding="utf-8").read().lower()
    manquants = [m for m in sorted(importes)
                 if equivalences.get(m, m).lower() not in declares]
    if manquants:
        return False, "absents de requirements.txt : " + ", ".join(manquants)
    return f"{len(importes)} bibliothèques"


# ─────────────────────────────────────────────────────────────────────
print()
if echecs:
    print("─" * 62)
    print(f" {ROUGE}✕  {len(echecs)} contrôle(s) en échec. Rien ne doit être envoyé.{NET}")
    print("─" * 62)
    print()
    sys.exit(1)

print("─" * 62)
print(f" {VERT}✓  Tout est bon. Le code peut partir.{NET}")
print("─" * 62)
print()
sys.exit(0)
